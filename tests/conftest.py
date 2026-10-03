"""Shared test fixtures: synthetic PDFs/DOCX built without network access."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from docx import Document


def make_png(path: Path, *, color: tuple[float, float, float] = (0.8, 0.2, 0.2)) -> Path:
    """Create a small solid-color PNG using PyMuPDF (no Pillow dependency)."""
    doc = pymupdf.open()
    page = doc.new_page(width=200, height=300)
    page.draw_rect(pymupdf.Rect(0, 0, 200, 300), color=color, fill=color)
    pixmap = page.get_pixmap()
    pixmap.save(str(path))
    doc.close()
    return path


def make_pdf(
    path: Path,
    *,
    pages: int = 3,
    header: str | None = "The Book",
    footer: str | None = "Confidential",
    body: str = "Body paragraph",
    metadata: dict[str, str] | None = None,
    image_path: Path | None = None,
    footer_image: Path | None = None,
) -> Path:
    """Create a synthetic multi-page PDF with an optional header/footer and image."""
    doc = pymupdf.open()
    for index in range(pages):
        page = doc.new_page()
        width = page.rect.width
        height = page.rect.height
        if header:
            page.insert_text((72, 40), header, fontsize=9)
        if footer:
            page.insert_text((72, height - 40), footer.format(n=index + 1), fontsize=9)
        if footer_image is not None:
            page.insert_image(
                pymupdf.Rect(20, height - 60, 160, height - 20),
                filename=str(footer_image),
            )
        page.insert_text((72, 120), f"{body} {index + 1}", fontsize=12)
        for line in range(6):
            page.insert_text(
                (72, 160 + line * 18),
                f"Line {line} of page {index + 1}.",
                fontsize=12,
            )
        if image_path is not None and index == 0:
            page.insert_image(pymupdf.Rect(72, 300, width - 72, 560), filename=str(image_path))
    if metadata:
        doc.set_metadata(metadata)
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def png_file(tmp_path: Path) -> Path:
    return make_png(tmp_path / "cover.png")


@pytest.fixture
def repeated_pdf(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "repeated.pdf")


@pytest.fixture
def page_number_pdf(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "numbers.pdf", header=None, footer="{n}")


@pytest.fixture
def clean_pdf_noise(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "clean.pdf", header=None, footer=None)


@pytest.fixture
def pdf_with_image(tmp_path: Path, png_file: Path) -> Path:
    return make_pdf(
        tmp_path / "illustrated.pdf",
        pages=2,
        header="The Book",
        footer="Confidential",
        image_path=png_file,
    )


@pytest.fixture
def docx_with_header_footer(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Body paragraph one.")
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Repeated header"
    section.footer.paragraphs[0].text = "Confidential"
    path = tmp_path / "noisy.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def docx_with_image(tmp_path: Path, png_file: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Before image.")
    doc.add_picture(str(png_file))
    doc.add_paragraph("After image.")
    path = tmp_path / "illustrated.docx"
    doc.save(str(path))
    return path

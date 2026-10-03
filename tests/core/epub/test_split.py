from __future__ import annotations

from pathlib import Path

from docx import Document

from pdf_to_epub.core.docx.clean import clean_document
from pdf_to_epub.core.docx.images import extract_images
from pdf_to_epub.core.epub.split import split_chapters


def test_splits_on_heading_styles() -> None:
    doc = Document()
    doc.add_heading("Chapter 1", level=1)
    doc.add_paragraph("First body.")
    doc.add_heading("Chapter 2", level=1)
    doc.add_paragraph("Second body.")

    chapters = split_chapters(doc, [], fallback_title="Book")
    assert [c.title for c in chapters] == ["Chapter 1", "Chapter 2"]
    assert "First body." in chapters[0].html
    assert "Second body." in chapters[1].html
    assert chapters[0].file_name == "chapter_0000.xhtml"


def test_splits_on_keyword_paragraph() -> None:
    doc = Document()
    doc.add_paragraph("Some front matter.")
    doc.add_paragraph("Prologue")
    doc.add_paragraph("Prologue body.")

    chapters = split_chapters(doc, [], fallback_title="Book")
    assert [c.title for c in chapters] == ["Book", "Prologue"]
    assert "Some front matter." in chapters[0].html
    assert "Prologue body." in chapters[1].html


def test_splits_on_guillemet_chapter_markers() -> None:
    doc = Document()
    doc.add_paragraph("«Prólogo»")
    doc.add_paragraph("Prologue body.")
    doc.add_paragraph("«Capítulo 1» El niño maldito")
    doc.add_paragraph("Chapter body.")

    chapters = split_chapters(doc, [], fallback_title="Book")
    assert [c.title for c in chapters] == ["«Prólogo»", "«Capítulo 1» El niño maldito"]
    assert "Prologue body." in chapters[0].html
    assert "Chapter body." in chapters[1].html


def test_toc_dot_leaders_are_not_chapters() -> None:
    doc = Document()
    doc.add_paragraph("Some body.")
    doc.add_paragraph("«Prólogo» .......................... 6")

    chapters = split_chapters(doc, [], fallback_title="Book")
    assert len(chapters) == 1


def test_splits_on_cap_markers() -> None:
    doc = Document()
    doc.add_paragraph("CAP 1: LA VISITA DOMICILIARIA")
    doc.add_paragraph("Cuerpo uno.")
    doc.add_paragraph("CAP. 2 Otra cosa")
    doc.add_paragraph("Cuerpo dos.")

    chapters = split_chapters(doc, [], fallback_title="Book")

    assert [c.title for c in chapters] == ["CAP 1: LA VISITA DOMICILIARIA", "CAP. 2 Otra cosa"]


def test_does_not_match_capital_word() -> None:
    doc = Document()
    doc.add_paragraph("capital de provincia es bonita")
    doc.add_paragraph("y sigue el texto.")

    chapters = split_chapters(doc, [], fallback_title="Book")

    assert len(chapters) == 1
    assert chapters[0].title == "Book"


def test_splits_on_cap_without_space() -> None:
    doc = Document()
    doc.add_paragraph("CAP1: LA VISITA DOMICILIARIA")
    doc.add_paragraph("Cuerpo uno.")
    doc.add_paragraph("CAP2: EL MES DE REFUERZO")
    doc.add_paragraph("CAP3: AKIRA ES...")
    doc.add_paragraph("Cuerpo tres.")

    chapters = split_chapters(doc, [], fallback_title="Book")

    assert [c.title for c in chapters] == [
        "CAP1: LA VISITA DOMICILIARIA",
        "CAP2: EL MES DE REFUERZO",
        "CAP3: AKIRA ES...",
    ]


def test_no_headings_yields_single_chapter() -> None:
    doc = Document()
    doc.add_paragraph("Only body.")
    chapters = split_chapters(doc, [], fallback_title="My Book")
    assert len(chapters) == 1
    assert chapters[0].title == "My Book"
    assert "Only body." in chapters[0].html


def test_embeds_image_reference(docx_with_image: Path) -> None:
    document = Document(str(docx_with_image))
    extraction = extract_images(document)
    chapters = split_chapters(document, extraction.assets, fallback_title="Book")
    html = "\n".join(chapter.html for chapter in chapters)
    assert 'img src="images/img_0000.png"' in html


def test_skip_rel_ids_omits_image_but_keeps_text(docx_with_image: Path) -> None:
    document = Document(str(docx_with_image))
    extraction = extract_images(document)
    rid = extraction.assets[0].rel_id
    chapters = split_chapters(
        document, extraction.assets, fallback_title="Book", skip_rel_ids={rid}
    )
    html = "\n".join(chapter.html for chapter in chapters)
    assert "img_0000" not in html
    assert "Before image." in html


def test_word_grid_table_is_repaired_before_rendering(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Before.")
    table = doc.add_table(rows=1, cols=4)
    for cell, word in zip(table.rows[0].cells, ["uno", "dos", "tres", "cuatro"], strict=True):
        cell.text = word
    path = tmp_path / "grid.docx"
    doc.save(str(path))

    clean_document(path)
    document = Document(str(path))
    chapters = split_chapters(document, [], fallback_title="Book")
    html = "\n".join(chapter.html for chapter in chapters)
    assert "<table" not in html
    assert "uno dos tres cuatro" in html
